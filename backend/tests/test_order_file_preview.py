from io import BytesIO

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from google import genai


def upload(client, name, content):
    return client.post("/api/v1/orders/file-preview", files={"file": (name, content)})


def test_text_excel_word_and_pdf_preview(client):
    response = upload(client, "order.txt", "10 κιλά ντομάτες".encode())
    assert response.status_code == 200
    assert response.json()["text"] == "10 κιλά ντομάτες"

    workbook = Workbook()
    workbook.active.append(["Product", "Qty"])
    workbook.active.append(["Tomatoes", 10])
    data = BytesIO()
    workbook.save(data)
    response = upload(client, "order.xlsx", data.getvalue())
    assert response.status_code == 200
    assert "Tomatoes | 10" in response.json()["text"]

    document = Document()
    document.add_paragraph("5 cartons cheese")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Olives"
    table.cell(0, 1).text = "3"
    data = BytesIO()
    document.save(data)
    response = upload(client, "order.docx", data.getvalue())
    assert response.status_code == 200
    assert "5 cartons cheese" in response.json()["text"]
    assert "Olives | 3" in response.json()["text"]

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    data = BytesIO()
    writer.write(data)
    response = upload(client, "scan.pdf", data.getvalue())
    assert response.status_code == 400
    assert "Gemini" in response.json()["detail"]


def test_excel_preview_preserves_blank_columns_for_ai_context(client):
    workbook = Workbook()
    workbook.active.append(["SKU", "Description", "M.M", "Qty", "Other", "SKU", "Description", "Qty"])
    workbook.active.append(["111315", "Paris", "TEM", None, "note", "180534", "Shoulder", "10+1"])
    data = BytesIO()
    workbook.save(data)
    response = upload(client, "mixed.xlsx", data.getvalue())
    assert response.status_code == 200
    assert "111315 | Paris | TEM |  | note | 180534 | Shoulder | 10+1" in response.json()["text"]


def test_upload_rejects_unsupported_or_oversize_content(client, monkeypatch):
    assert upload(client, "old.doc", b"legacy").status_code == 400
    assert upload(client, "photo.jpg", b"not a jpeg").status_code == 400
    assert upload(client, "order.pdf", b"not a pdf").status_code == 400
    monkeypatch.setattr("backend.app.api.v1.orders.settings.MAX_ORDER_UPLOAD_SIZE_BYTES", 5)
    assert upload(client, "order.txt", b"123456").status_code == 413


def test_handwritten_photo_uses_gemini_transcription_and_flags_unclear_text(client, monkeypatch):
    calls = []

    class Overloaded(Exception):
        code = 503

    class FakeModels:
        async def generate_content(self, **kwargs):
            calls.append(kwargs)
            if len(calls) == 1:
                raise Overloaded()
            return type("Response", (), {"text": "10 [UNCLEAR] olives"})()

    class FakeAsyncClient:
        models = FakeModels()

        async def aclose(self):
            pass

    class FakeClient:
        aio = FakeAsyncClient()

        def close(self):
            pass

    monkeypatch.setattr(genai, "Client", lambda **kwargs: FakeClient())
    async def no_delay(_seconds):
        pass
    monkeypatch.setattr("backend.app.services.order_file_service.asyncio.sleep", no_delay)
    monkeypatch.setattr("backend.app.services.order_file_service.settings.AI_PROVIDER", "gemini")
    monkeypatch.setattr("backend.app.services.order_file_service.settings.GEMINI_API_KEY", "test-key")
    response = upload(client, "note.jpg", b"\xff\xd8\xfffake-image")
    assert response.status_code == 200
    assert response.json()["text"] == "10 [UNCLEAR] olives"
    assert response.json()["method"] == "ocr"
    assert len(calls) == 2
    assert "handwritten" in calls[0]["contents"][0]
    assert calls[0]["contents"][1].inline_data.mime_type == "image/jpeg"
