import pytest
import io
import csv
import json
import openpyxl
from sqlalchemy.orm import Session

from backend.app.models.company import Company
from backend.app.models.business_settings import CompanyBusinessSettings
from backend.app.models.customer import Customer
from backend.app.models.product import Product, Packaging
from backend.app.models.order import Order
from backend.app.schemas.workflow import OrderStatus
from backend.app.schemas.matching import LineMatchResult, MatchedProductInfo, ConfidenceResult, MatchDecision
from backend.app.schemas.export import (
    ExportProfileCreate,
    ExportProfileUpdate,
    ExportFieldMappingCreate,
    MappingType,
)
from backend.app.services.order_workflow_service import OrderWorkflowService
from backend.app.services.export_profile_service import (
    ExportProfileService,
    ExportProfileValidationError,
)
from backend.app.services.export_engine import (
    ExportEngine,
    OrderExportError,
)


def setup_export_data(db_session: Session, empty_master_fields: bool = False):
    company = Company(name="Hellas Food Logistics")
    db_session.add(company)
    db_session.flush()
    db_session.add(CompanyBusinessSettings(company_id=company.id, unitless_order_behavior="piece"))

    customer = Customer(
        company_id=company.id,
        customer_code="CUST-100",
        customer_name="Grand Hotel Athens",
        email="orders@grandhotel.gr",
        phone=None if empty_master_fields else "2100000000"
    )
    prod1 = Product(
        company_id=company.id,
        sku="SKU-7843",
        description="Γαλοπούλα Καπνιστή 1kg",
        barcode=None if empty_master_fields else "5201234567890",
        unit="piece",
        active=True
    )
    prod2 = Product(
        company_id=company.id,
        sku="SKU-100",
        description="Σαλάμι Μπύρας 300g",
        barcode="5209876543210",
        unit="piece",
        active=True
    )
    db_session.add_all([customer, prod1, prod2])
    db_session.commit()

    # Create approved order
    line1 = LineMatchResult(
        line_number=1,
        original_text="10 γαλοπουλες",
        product_phrase="γαλοπουλες",
        quantity=10.0,
        unit="piece",
        best_match=MatchedProductInfo(product_id=prod1.id, sku=prod1.sku, description=prod1.description, unit=prod1.unit),
        confidence=ConfidenceResult(score=0.98, decision=MatchDecision.AUTO_ACCEPT)
    )
    line2 = LineMatchResult(
        line_number=2,
        original_text="5 σαλαμια",
        product_phrase="σαλαμια",
        quantity=5.0,
        unit="piece",
        best_match=MatchedProductInfo(product_id=prod2.id, sku=prod2.sku, description=prod2.description, unit=prod2.unit),
        confidence=ConfidenceResult(score=0.96, decision=MatchDecision.AUTO_ACCEPT)
    )

    order = OrderWorkflowService.create_order_from_match(
        db=db_session,
        company_id=company.id,
        customer_id=customer.id,
        raw_input="10 γαλοπουλες\n5 σαλαμια",
        lines=[line1, line2]
    )
    OrderWorkflowService.approve_order(db=db_session, order_id=order.id)

    return {
        "company": company,
        "customer": customer,
        "prod1": prod1,
        "prod2": prod2,
        "order": order
    }


def test_export_profile_validation_rules(db_session):
    data = setup_export_data(db_session)
    comp_id = data["company"].id

    # 1. Empty mappings rejected
    with pytest.raises(ExportProfileValidationError, match="at least one column mapping"):
        ExportProfileService.validate_profile_mappings([])

    # 2. Duplicate column_order rejected
    with pytest.raises(ExportProfileValidationError, match="Duplicate column_order"):
        ExportProfileService.create_profile(
            db=db_session,
            payload=ExportProfileCreate(
                company_id=comp_id,
                name="Dup Order",
                mappings=[
                    ExportFieldMappingCreate(column_order=1, output_column_name="A", source_field="line.sku"),
                    ExportFieldMappingCreate(column_order=1, output_column_name="B", source_field="line.quantity")
                ]
            )
        )


    # 3. Invalid source_field rejected
    with pytest.raises(ExportProfileValidationError, match="Invalid source_field"):
        ExportProfileService.create_profile(
            db=db_session,
            payload=ExportProfileCreate(
                company_id=comp_id,
                name="Invalid Field",
                mappings=[
                    ExportFieldMappingCreate(column_order=1, output_column_name="X", source_field="hacker.secret_token")
                ]
            )
        )

    # 4. Constant without constant_value rejected
    with pytest.raises(ExportProfileValidationError, match="no constant_value specified"):
        ExportProfileService.create_profile(
            db=db_session,
            payload=ExportProfileCreate(
                company_id=comp_id,
                name="Bad Constant",
                mappings=[
                    ExportFieldMappingCreate(column_order=1, output_column_name="WH", mapping_type=MappingType.CONSTANT, constant_value=None)
                ]
            )
        )


def test_order_sheet_bonus_marker_can_be_cleared(db_session):
    company = Company(name="Marker company")
    db_session.add(company)
    db_session.commit()
    profile = ExportProfileService.create_profile(db_session, ExportProfileCreate(
        company_id=company.id, name="Gift sheet", format="order_sheet", mappings=[],
        bonus_separate_row=True, bonus_marker="Α",
    ))
    updated = ExportProfileService.update_profile(db_session, profile.id, ExportProfileUpdate(
        bonus_separate_row=False, bonus_marker=None,
    ))
    assert updated.bonus_separate_row is False
    assert updated.bonus_marker is None


def test_export_xlsx(db_session):
    data = setup_export_data(db_session)
    comp = data["company"]
    order = data["order"]

    profile = ExportProfileService.create_profile(
        db=db_session,
        payload=ExportProfileCreate(
            company_id=comp.id,
            name="Warehouse Excel",
            format="xlsx",
            include_header=True,
            mappings=[
                ExportFieldMappingCreate(column_order=1, output_column_name="CUSTOMER", source_field="customer.customer_code"),
                ExportFieldMappingCreate(column_order=2, output_column_name="ITEM", source_field="line.sku"),
                ExportFieldMappingCreate(column_order=3, output_column_name="QTY", source_field="line.quantity"),
                ExportFieldMappingCreate(column_order=4, output_column_name="UNIT", source_field="line.unit"),
                ExportFieldMappingCreate(column_order=5, output_column_name="WAREHOUSE", mapping_type=MappingType.CONSTANT, constant_value="01")
            ]
        )
    )

    file_bytes, media_type, filename = ExportEngine.export_order(
        db=db_session,
        order_id=order.id,
        profile_id=profile.id
    )

    assert media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert filename.endswith(".xlsx")
    assert f"ordermind_order_{order.id}_warehouse_excel" in filename

    # Inspect openpyxl workbook
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    ws = wb.active
    assert ws.title == "Order"

    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 3  # Header + 2 lines

    # Header
    assert rows[0] == ("CUSTOMER", "ITEM", "QTY", "UNIT", "WAREHOUSE")
    # Row 1
    assert rows[1] == ("CUST-100", "SKU-7843", 10.0, "piece", "01")
    # Row 2
    assert rows[2] == ("CUST-100", "SKU-100", 5.0, "piece", "01")

    # Verify order audit
    db_session.refresh(order)
    assert order.status == OrderStatus.EXPORTED.value
    assert order.exported_at is not None
    assert order.last_export_profile_id == profile.id


def test_four_column_order_sheet_converts_cases_and_separates_bonus_from_snapshot(db_session, client):
    company = Company(name="Order sheet company")
    db_session.add(company)
    db_session.flush()
    db_session.add(CompanyBusinessSettings(company_id=company.id, bonus_enabled=True, bonus_expression_mode="paid_plus_bonus", allow_packaging_conversion=True))
    profile = ExportProfileService.create_profile(db_session, ExportProfileCreate(
        company_id=company.id, name="Four columns", format="order_sheet", mappings=[], include_header=False,
        bonus_separate_row=True, bonus_marker="Α", quantity_output_unit="piece",
        convert_case_using_pieces_per_case=True,
    ))
    source_profile = ExportProfileService.create_profile(db_session, ExportProfileCreate(
        company_id=company.id, name="Source units", format="order_sheet", mappings=[], include_header=False,
        bonus_separate_row=True, bonus_marker="FREE", quantity_output_unit="source",
    ))
    no_bonus_profile = ExportProfileService.create_profile(db_session, ExportProfileCreate(
        company_id=company.id, name="No bonus row", format="order_sheet", mappings=[], include_header=False,
    ))
    customer = Customer(company_id=company.id, customer_code="SHEET", customer_name="Buyer")
    product = Product(company_id=company.id, sku="180534", description="ΩΜΟΠΛΑΤΗ", unit="piece")
    db_session.add_all([customer, product])
    db_session.flush()
    package = Packaging(company_id=company.id, product_id=product.id, package_type="case", pieces_per_case=12, unit="piece")
    db_session.add(package)
    db_session.commit()
    match = LineMatchResult(
        line_number=1, original_text="180534 10 κιβώτια + 1 δώρο", product_phrase="180534",
        quantity=10, unit="case", raw_unit="κιβώτια", unit_explicit=True,
        quantity_text="10 κιβώτια + 1 δώρο", bonus_quantity=1,
        best_match=MatchedProductInfo(product_id=product.id, sku=product.sku, description=product.description, unit="piece"),
        matched_packaging_id=package.id,
        confidence=ConfidenceResult(score=0.98, decision=MatchDecision.AUTO_ACCEPT),
    )
    order = OrderWorkflowService.create_order_from_match(
        db_session, company.id, customer.id, match.original_text, [match]
    )
    pending = client.get(f"/api/v1/orders/{order.id}?profile_id={profile.id}")
    assert pending.status_code == 200, pending.text
    assert pending.json()["lines"][0]["order_sheet_paid_quantity"] == 120
    assert pending.json()["lines"][0]["order_sheet_bonus_quantity"] == 12
    assert pending.json()["lines"][0]["order_sheet_unit"] == "piece"
    OrderWorkflowService.approve_order(db_session, order.id)
    package.pieces_per_case = 24
    db_session.commit()
    approved = client.get(f"/api/v1/orders/{order.id}?profile_id={profile.id}")
    assert approved.status_code == 200, approved.text
    assert approved.json()["lines"][0]["order_sheet_paid_quantity"] == 120
    assert approved.json()["lines"][0]["order_sheet_bonus_quantity"] == 12
    unsupported_preview = client.get(f"/api/v1/orders/{order.id}?profile_id={no_bonus_profile.id}")
    assert unsupported_preview.status_code == 200
    assert unsupported_preview.json()["lines"][0]["order_sheet_paid_quantity"] is None
    assert "does not define how to export bonus goods" in unsupported_preview.json()["lines"][0]["order_sheet_conversion_error"]
    with pytest.raises(OrderExportError, match="does not define how to export bonus goods"):
        ExportEngine.export_order(db_session, order.id, no_bonus_profile.id)
    content, _, _ = ExportEngine.export_order(db_session, order.id, profile.id)
    rows = list(openpyxl.load_workbook(io.BytesIO(content)).active.values)
    assert rows == [
        ("180534", "ΩΜΟΠΛΑΤΗ", None, 120),
        ("180534", "ΩΜΟΠΛΑΤΗ", "Α", 12),
    ]
    source_content, _, _ = ExportEngine.export_order(db_session, order.id, source_profile.id)
    assert list(openpyxl.load_workbook(io.BytesIO(source_content)).active.values) == [
        ("180534", "ΩΜΟΠΛΑΤΗ", None, 10),
        ("180534", "ΩΜΟΠΛΑΤΗ", "FREE", 1),
    ]
    profile.bonus_marker = "CHANGED"
    profile.quantity_output_unit = "source"
    profile.convert_case_using_pieces_per_case = False
    db_session.commit()
    frozen_content, _, _ = ExportEngine.export_order(db_session, order.id, profile.id)
    assert list(openpyxl.load_workbook(io.BytesIO(frozen_content)).active.values) == rows
    frozen_preview = client.get(f"/api/v1/orders/{order.id}?profile_id={profile.id}")
    assert frozen_preview.json()["lines"][0]["order_sheet_bonus_marker"] == "Α"


def test_create_from_match_preserves_grounded_bonus(client, db_session):
    company = Company(name="Bonus API company")
    db_session.add(company)
    db_session.flush()
    db_session.add(CompanyBusinessSettings(company_id=company.id, bonus_enabled=True, bonus_expression_mode="paid_plus_bonus"))
    customer = Customer(company_id=company.id, customer_code="BONUS", customer_name="Bonus buyer")
    product = Product(company_id=company.id, sku="131382", description="Frank", unit="piece")
    db_session.add_all([customer, product])
    db_session.commit()
    text = "131382 Frank 56+6"
    response = client.post("/api/v1/orders/create-from-match", json={
        "company_id": company.id, "customer_id": customer.id, "raw_input": text,
        "items": [{
            "line_number": 1, "original_text": text, "product_phrase": "131382",
            "quantity": 56, "quantity_text": "56+6", "bonus_quantity": 6, "unit": "piece",
        }],
    })
    assert response.status_code == 201, response.text
    assert response.json()["lines"][0]["bonus_quantity"] == 6


def test_order_sheet_rejects_unproven_conversion_and_keeps_neutral_source_unit(db_session):
    from backend.app.services.export_engine import convert_order_sheet_quantities

    with pytest.raises(OrderExportError, match="not enabled"):
        convert_order_sheet_quantities(10, 1, "case", 12, "piece", True, False)
    with pytest.raises(OrderExportError, match="valid case packaging ratio"):
        convert_order_sheet_quantities(10, 1, "case", None, "piece", True, True)
    with pytest.raises(OrderExportError, match="cannot be converted"):
        convert_order_sheet_quantities(10, 0, "kg", None, "piece", True, True)

    data = setup_export_data(db_session)
    company = data["company"]
    profile = ExportProfileService.create_profile(db_session, ExportProfileCreate(
        company_id=company.id, name="Neutral sheet", format="order_sheet",
        mappings=[], include_header=False,
    ))
    # A no-bonus order remains usable with neutral source-unit output.
    content, _, _ = ExportEngine.export_order(db_session, data["order"].id, profile.id)
    assert len(list(openpyxl.load_workbook(io.BytesIO(content)).active.values)) == 2


def test_reexport_uses_approved_snapshot_including_empty_master_fields(db_session):
    data = setup_export_data(db_session, empty_master_fields=True)
    order = data["order"]
    customer = data["customer"]
    product = data["prod1"]
    profile = ExportProfileService.create_profile(
        db=db_session,
        payload=ExportProfileCreate(
            company_id=data["company"].id,
            name="Snapshot JSON",
            format="json",
            mappings=[
                ExportFieldMappingCreate(column_order=1, output_column_name="customer", source_field="customer.customer_code"),
                ExportFieldMappingCreate(column_order=2, output_column_name="phone", source_field="customer.phone"),
                ExportFieldMappingCreate(column_order=3, output_column_name="sku", source_field="product.sku"),
                ExportFieldMappingCreate(column_order=4, output_column_name="description", source_field="product.description"),
                ExportFieldMappingCreate(column_order=5, output_column_name="barcode", source_field="product.barcode"),
            ],
        ),
    )
    first, _, _ = ExportEngine.export_order(db_session, order.id, profile.id)
    customer.customer_code = "CHANGED"
    customer.phone = "999"
    product.sku = "CHANGED-SKU"
    product.description = "Changed description"
    product.barcode = "999999"
    db_session.commit()
    second, _, _ = ExportEngine.export_order(db_session, order.id, profile.id)
    assert json.loads(first) == json.loads(second)
    row = json.loads(second)[0]
    assert row == {"customer": "CUST-100", "phone": "", "sku": "SKU-7843", "description": "Γαλοπούλα Καπνιστή 1kg", "barcode": ""}
    assert len(order.export_records) == 2


def test_export_csv(db_session):
    data = setup_export_data(db_session)
    comp = data["company"]
    order = data["order"]

    profile = ExportProfileService.create_profile(
        db=db_session,
        payload=ExportProfileCreate(
            company_id=comp.id,
            name="SoftOne CSV",
            format="csv",
            delimiter=";",
            include_header=True,
            encoding="utf-8-sig",
            mappings=[
                ExportFieldMappingCreate(column_order=1, output_column_name="CODE", source_field="customer.customer_code"),
                ExportFieldMappingCreate(column_order=2, output_column_name="SKU", source_field="line.sku"),
                ExportFieldMappingCreate(column_order=3, output_column_name="QUANTITY", source_field="line.quantity"),
                ExportFieldMappingCreate(column_order=4, output_column_name="SERIES", mapping_type=MappingType.CONSTANT, constant_value="PAR")
            ]
        )
    )

    file_bytes, media_type, filename = ExportEngine.export_order(
        db=db_session,
        order_id=order.id,
        profile_id=profile.id
    )

    assert "text/csv" in media_type
    assert filename.endswith(".csv")

    decoded_text = file_bytes.decode("utf-8-sig")
    reader = list(csv.reader(io.StringIO(decoded_text), delimiter=";"))

    assert len(reader) == 3
    assert reader[0] == ["CODE", "SKU", "QUANTITY", "SERIES"]
    assert reader[1] == ["CUST-100", "SKU-7843", "10.0", "PAR"]
    assert reader[2] == ["CUST-100", "SKU-100", "5.0", "PAR"]


def test_export_json(db_session):
    data = setup_export_data(db_session)
    comp = data["company"]
    order = data["order"]

    profile = ExportProfileService.create_profile(
        db=db_session,
        payload=ExportProfileCreate(
            company_id=comp.id,
            name="ERP JSON",
            format="json",
            mappings=[
                ExportFieldMappingCreate(column_order=1, output_column_name="CUSTOMER", source_field="customer.customer_code"),
                ExportFieldMappingCreate(column_order=2, output_column_name="SKU", source_field="line.sku"),
                ExportFieldMappingCreate(column_order=3, output_column_name="QTY", source_field="line.quantity"),
                ExportFieldMappingCreate(column_order=4, output_column_name="TAG", mapping_type=MappingType.CONSTANT, constant_value="B2B")
            ]
        )
    )

    file_bytes, media_type, filename = ExportEngine.export_order(
        db=db_session,
        order_id=order.id,
        profile_id=profile.id
    )

    assert media_type == "application/json; charset=utf-8"
    assert filename.endswith(".json")

    data_json = json.loads(file_bytes.decode("utf-8"))
    assert len(data_json) == 2
    assert data_json[0] == {
        "CUSTOMER": "CUST-100",
        "SKU": "SKU-7843",
        "QTY": 10.0,
        "TAG": "B2B"
    }
    assert data_json[1] == {
        "CUSTOMER": "CUST-100",
        "SKU": "SKU-100",
        "QTY": 5.0,
        "TAG": "B2B"
    }


def test_unapproved_order_export_blocked(db_session):
    data = setup_export_data(db_session)
    comp = data["company"]
    cust = data["customer"]

    # Create unapproved pending_review order
    unapproved_order = Order(
        company_id=comp.id,
        customer_id=cust.id,
        order_number="ORD-PENDING-001",
        status=OrderStatus.PENDING_REVIEW.value,
        raw_input="raw"
    )
    db_session.add(unapproved_order)
    db_session.commit()

    profile = ExportProfileService.create_profile(
        db=db_session,
        payload=ExportProfileCreate(
            company_id=comp.id,
            name="Test Profile",
            mappings=[
                ExportFieldMappingCreate(column_order=1, output_column_name="SKU", source_field="line.sku")
            ]
        )
    )

    # 1. Final export must be blocked
    with pytest.raises(OrderExportError, match="cannot be exported because its status is 'pending_review'"):
        ExportEngine.export_order(db=db_session, order_id=unapproved_order.id, profile_id=profile.id)

    # 2. Preview export is allowed and does NOT mutate order status or exported_at
    file_bytes, _, _ = ExportEngine.export_order(
        db=db_session,
        order_id=unapproved_order.id,
        profile_id=profile.id,
        preview=True
    )
    assert len(file_bytes) > 0
    db_session.refresh(unapproved_order)
    assert unapproved_order.status == OrderStatus.PENDING_REVIEW.value
    assert unapproved_order.exported_at is None


def test_formula_injection_prevention(db_session):
    """
    Ensure values starting with '=', '+', '-', '@', '\t', '\r' are escaped with single quote.
    """
    data = setup_export_data(db_session)
    comp = data["company"]
    cust = data["customer"]
    prod1 = data["prod1"]
    prod1.description = "+SUM(A1:A10)"
    db_session.commit()

    # Malicious text input
    line_malicious = LineMatchResult(
        line_number=1,
        original_text='=HYPERLINK("http://attacker.com")',
        product_phrase='=HYPERLINK("http://attacker.com")',
        quantity=1.0,
        unit="piece",
        best_match=MatchedProductInfo(product_id=prod1.id, sku="SKU-7843", description=prod1.description, unit="piece"),
        confidence=ConfidenceResult(score=0.99, decision=MatchDecision.AUTO_ACCEPT)
    )

    order = OrderWorkflowService.create_order_from_match(
        db=db_session,
        company_id=comp.id,
        customer_id=cust.id,
        raw_input="malicious input",
        lines=[line_malicious]
    )
    OrderWorkflowService.approve_order(db=db_session, order_id=order.id)

    profile = ExportProfileService.create_profile(
        db=db_session,
        payload=ExportProfileCreate(
            company_id=comp.id,
            name="Safe Profile",
            format="csv",
            delimiter=",",
            mappings=[
                ExportFieldMappingCreate(column_order=1, output_column_name="PHRASE", source_field="line.product_phrase"),
                ExportFieldMappingCreate(column_order=2, output_column_name="DESC", source_field="line.description"),
                ExportFieldMappingCreate(column_order=3, output_column_name="CONSTANT_EVIL", mapping_type=MappingType.CONSTANT, constant_value="@EXEC('calc.exe')")
            ]
        )
    )

    file_bytes, _, _ = ExportEngine.export_order(db=db_session, order_id=order.id, profile_id=profile.id)
    reader = list(csv.reader(io.StringIO(file_bytes.decode("utf-8-sig"))))

    data_row = reader[1]
    # Verify values are escaped with leading single quote
    assert data_row[0] == "'=HYPERLINK(\"http://attacker.com\")"
    assert data_row[1] == "'+SUM(A1:A10)"
    assert data_row[2] == "'@EXEC('calc.exe')"


def test_company_isolation_on_export(db_session):
    data = setup_export_data(db_session)
    order = data["order"]

    other_comp = Company(name="Rival Corp")
    db_session.add(other_comp)
    db_session.flush()

    foreign_profile = ExportProfileService.create_profile(
        db=db_session,
        payload=ExportProfileCreate(
            company_id=other_comp.id,
            name="Foreign Profile",
            mappings=[
                ExportFieldMappingCreate(column_order=1, output_column_name="SKU", source_field="line.sku")
            ]
        )
    )

    with pytest.raises(ValueError, match="belongs to company"):
        ExportEngine.export_order(db=db_session, order_id=order.id, profile_id=foreign_profile.id)


def test_export_header_formula_injection_sanitization(db_session):
    """
    Verify that export headers starting with formula triggers (=, +, -, @)
    are sanitized with a leading quote in XLSX and CSV output,
    while leaving the stored mapping output_column_name intact in the database.
    """
    data = setup_export_data(db_session)
    comp = data["company"]
    order = data["order"]

    profile = ExportProfileService.create_profile(
        db=db_session,
        payload=ExportProfileCreate(
            company_id=comp.id,
            name="Formula Header Profile",
            format="csv",
            delimiter=";",
            mappings=[
                ExportFieldMappingCreate(column_order=1, output_column_name="=CMD|' /C calc'!A0", source_field="line.sku"),
                ExportFieldMappingCreate(column_order=2, output_column_name="+SUM_COL", source_field="line.quantity"),
                ExportFieldMappingCreate(column_order=3, output_column_name="-DIFF_COL", source_field="line.unit"),
                ExportFieldMappingCreate(column_order=4, output_column_name="@USER_COL", source_field="customer.customer_name"),
            ]
        )
    )

    # 1. Test CSV export headers
    file_bytes, media_type, filename = ExportEngine.export_order(
        db=db_session, order_id=order.id, profile_id=profile.id, preview=True
    )
    csv_text = file_bytes.decode("utf-8-sig")
    reader = list(csv.reader(io.StringIO(csv_text), delimiter=";"))
    csv_headers = reader[0]
    assert csv_headers[0] == "'=CMD|' /C calc'!A0"
    assert csv_headers[1] == "'+SUM_COL"
    assert csv_headers[2] == "'-DIFF_COL"
    assert csv_headers[3] == "'@USER_COL"

    # 2. Test XLSX export headers
    profile.format = "xlsx"
    db_session.commit()

    xlsx_bytes, _, _ = ExportEngine.export_order(
        db=db_session, order_id=order.id, profile_id=profile.id, preview=True
    )
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    xlsx_headers = [cell.value for cell in ws[1]]
    assert xlsx_headers[0] == "'=CMD|' /C calc'!A0"
    assert xlsx_headers[1] == "'+SUM_COL"
    assert xlsx_headers[2] == "'-DIFF_COL"
    assert xlsx_headers[3] == "'@USER_COL"

    # 3. Database stored mappings must remain UNTOUCHED
    db_session.refresh(profile)
    stored_names = [m.output_column_name for m in profile.field_mappings]
    assert "=CMD|' /C calc'!A0" in stored_names
    assert "+SUM_COL" in stored_names
    assert "-DIFF_COL" in stored_names
    assert "@USER_COL" in stored_names
