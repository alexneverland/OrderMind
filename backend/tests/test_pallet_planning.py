"""Physical pallet invariants and profile-scoped output."""
import io
import copy
import zipfile
from decimal import Decimal

import openpyxl
import pytest
from pydantic import ValidationError

from backend.app.models.company import Company
from backend.app.models.customer import Customer
from backend.app.models.product import Product
from backend.app.models.order import Order, OrderLine
from backend.app.models.export import ExportProfile
from backend.app.services.order_workflow_service import OrderWorkflowService
from backend.app.services.export_engine import ExportEngine
from backend.app.schemas.export import ExportProfileCreate
from backend.app.schemas.pallet import PalletConfig
from backend.app.services.export_engine import OrderExportError
from backend.app.services.export_profile_service import ExportProfileService, ExportProfileValidationError
from backend.app.services.pallet_planner import plan_pallets
from backend.app.services.pallet_renderer import render_pallet_plan
from backend.app.services.master_data_service import MasterDataService
from backend.app.services.rules_assistant import _mock_analysis

POLICY = {"bonus_separate_row": True, "bonus_marker": "Α", "quantity_output_unit": "source",
          "convert_case_using_pieces_per_case": False, "include_header": False}
BUSINESS = {"allow_packaging_conversion": False}


def line(number, sku, paid, bonus=0, unit="piece", **weight):
    return {"line_id": number, "line_number": number, "product_id": number, "sku": sku,
            "description": f"Product {sku}", "quantity": paid, "bonus_quantity": bonus,
            "unit": unit, "pieces_per_case": weight.get("pieces_per_case"),
            "kg_per_piece": weight.get("kg_per_piece"), "kg_per_case": weight.get("kg_per_case")}


def config(**automatic):
    return PalletConfig.model_validate({"enabled": True, "automatic_pallets": automatic})


def sums(plan, line_id):
    chunks = [item for pallet in plan["pallets"] for item in pallet["items"] if item["source_order_line_id"] == line_id]
    return sum(Decimal(item["paid_quantity"]) for item in chunks), sum(Decimal(item["bonus_quantity"]) for item in chunks)


def test_piece_and_bonus_weight_split_conserves_both():
    plan = plan_pallets([line(1, "A", 12, 2, kg_per_piece="1.5")], config(max_weight_kg=10), POLICY, BUSINESS)
    assert len(plan["pallets"]) == 3
    assert sums(plan, 1) == (Decimal(12), Decimal(2))
    assert all(Decimal(pallet["total_weight_kg"]) <= 10 for pallet in plan["pallets"])
    assert all(Decimal(item["paid_quantity"]) == Decimal(item["paid_quantity"]).to_integral_value()
               and Decimal(item["bonus_quantity"]) == Decimal(item["bonus_quantity"]).to_integral_value()
               for pallet in plan["pallets"] for item in pallet["items"])


def test_kg_and_case_mass_and_untrusted_weight():
    plan = plan_pallets([line(1, "KG", 500, unit="kg")], config(max_weight_kg=360), POLICY, BUSINESS)
    assert [p["total_weight_kg"] for p in plan["pallets"]] == ["360", "140"]
    case = line(2, "CASE", 3, 1, unit="case", pieces_per_case=12, kg_per_piece="0.5")
    plan = plan_pallets([case], config(max_weight_kg=24), POLICY, BUSINESS)
    assert plan["pallets"][0]["total_weight_kg"] == "24.0"
    with pytest.raises(OrderExportError, match="SKU CASE.*no frozen trusted"):
        plan_pallets([line(2, "CASE", 1, unit="case")], config(max_weight_kg=10), POLICY, BUSINESS)


def test_row_modes_and_dedicated_groups():
    items = [line(1, "A", 1, 1, kg_per_piece=1), line(2, "B", 1, kg_per_piece=1), line(3, "C", 1, kg_per_piece=1)]
    cfg = PalletConfig.model_validate({"enabled": True, "dedicated_groups": [{"name": "Together", "product_ids": [1, 2]}],
        "automatic_pallets": {"max_rows": 1, "max_weight_kg": 1}})
    plan = plan_pallets(items, cfg, POLICY, BUSINESS)
    assert len(plan["pallets"]) == 2
    assert plan["pallets"][0]["type"] == "dedicated"
    assert plan["pallets"][0]["row_count"] == 3
    assert len(plan["pallets"][0]["warnings"]) == 2
    cfg.automatic_pallets.row_count_mode = "logical_product_lines"
    plan = plan_pallets(items, cfg, POLICY, BUSINESS)
    assert plan["pallets"][0]["row_count"] == 2


@pytest.mark.parametrize("layout", ["single_sheet_sections", "multi_sheet_workbook", "separate_workbook_per_pallet"])
def test_layouts_preserve_product_rows(layout):
    cfg = PalletConfig.model_validate({"enabled": True, "automatic_pallets": {"max_rows": 2},
        "output": {"layout": layout, "show_pallet_title": True, "repeat_headers": True, "blank_rows_between_pallets": 1}})
    plan = plan_pallets([line(1, "A", 1, 1), line(2, "B", 1)], cfg, POLICY, BUSINESS)
    content, media, filename = render_pallet_plan(10, plan, cfg, False)
    if layout == "separate_workbook_per_pallet":
        assert media == "application/zip" and filename.endswith(".zip")
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            assert len(zf.namelist()) == 2
            books = [openpyxl.load_workbook(io.BytesIO(zf.read(name))) for name in zf.namelist()]
    else:
        books = [openpyxl.load_workbook(io.BytesIO(content))]
    actual = [row for book in books for sheet in book.worksheets for row in sheet.values
              if row and row[0] in {"A", "B"}]
    assert len(actual) == 3
    if layout == "multi_sheet_workbook":
        assert len(books[0].worksheets) == 2
        assert [sheet.title for sheet in books[0].worksheets] == ["Pallet 1", "Pallet 2"]
    elif layout == "single_sheet_sections":
        values = list(books[0].active.values)
        assert [row[0] for row in values if row[0] and str(row[0]).startswith("PALLET")] == ["PALLET 1", "PALLET 2"]
        assert sum(row[0] == "SKU" for row in values) == 2
    else:
        assert all(len(book.worksheets) == 1 for book in books)


def test_profile_validation_blocks_cross_company_membership(db_session):
    a, b = Company(name="A"), Company(name="B")
    db_session.add_all([a, b]); db_session.flush()
    product = Product(company_id=b.id, sku="B", description="B")
    db_session.add(product); db_session.flush()
    cfg = PalletConfig.model_validate({"enabled": True, "dedicated_groups": [{"name": "Bad", "product_ids": [product.id]}],
        "automatic_pallets": {"max_rows": 10}})
    with pytest.raises(ExportProfileValidationError, match="must belong"):
        ExportProfileService.validate_palletization(db_session, a.id, "order_sheet", cfg)


def test_policy_rejects_duplicate_membership_and_missing_capacity():
    with pytest.raises(ValidationError, match="only one dedicated"):
        PalletConfig.model_validate({"enabled": True, "dedicated_groups": [
            {"name": "One", "product_ids": [1]}, {"name": "Two", "product_ids": [1]}],
            "automatic_pallets": {"max_rows": 2}})
    with pytest.raises(ValidationError, match="At least one"):
        PalletConfig.model_validate({"enabled": True})


def test_row_capacity_counts_actual_bonus_rows_and_keeps_pair_together():
    items = [line(number, f"SKU-{number}", 1) for number in range(1, 16)]
    items.append(line(16, "GIFT", 5, 1))
    cfg = config(max_rows=16)
    plan = plan_pallets(items, cfg, POLICY, BUSINESS)
    assert [pallet["row_count"] for pallet in plan["pallets"]] == [15, 2]
    assert [item["sku"] for item in plan["pallets"][1]["items"]] == ["GIFT"]
    cfg.automatic_pallets.row_count_mode = "logical_product_lines"
    logical = plan_pallets(items, cfg, POLICY, BUSINESS)
    assert len(logical["pallets"]) == 1 and logical["pallets"][0]["row_count"] == 16


def test_single_physical_unit_too_heavy_fails_with_sku():
    with pytest.raises(OrderExportError, match="SKU HUGE.*one physical unit"):
        plan_pallets([line(1, "HUGE", 1, kg_per_piece=11)], config(max_weight_kg=10), POLICY, BUSINESS)


def test_pathological_plan_has_bounded_output():
    with pytest.raises(OrderExportError, match="exceeds the supported 2000 pallets"):
        plan_pallets([line(1, "BULK", 2001, unit="kg")], config(max_weight_kg=1), POLICY, BUSINESS)


def test_approval_freezes_weight_and_policy_and_audits_zip(client, db_session):
    company = Company(name="Frozen")
    db_session.add(company); db_session.flush()
    customer = Customer(company_id=company.id, customer_code="C", customer_name="Buyer")
    product = Product(company_id=company.id, sku="A", description="A", unit="piece", kg_per_piece=2)
    db_session.add_all([customer, product]); db_session.flush()
    cfg = PalletConfig.model_validate({"enabled": True, "automatic_pallets": {"max_weight_kg": 6},
        "output": {"layout": "separate_workbook_per_pallet"}})
    profile = ExportProfileService.create_profile(db_session, ExportProfileCreate(
        company_id=company.id, name="Pallet ZIP", format="order_sheet", include_header=False,
        palletization=cfg))
    order = Order(company_id=company.id, customer_id=customer.id, order_number="O-1", raw_input="5 A", status="pending_review")
    db_session.add(order); db_session.flush()
    order.lines.append(OrderLine(company_id=company.id, line_number=1, original_text="5 A", product_phrase="A",
        requested_quantity=5, requested_unit="piece", final_quantity=5, final_unit="piece", matched_product_id=product.id,
        final_sku="A", confidence_score=1, status="confirmed"))
    db_session.commit()
    OrderWorkflowService.approve_order(db_session, order.id)
    assert client.get(f"/api/v1/orders/{order.id}").json()["pallet_profile_ids"] == [profile.id]
    assert Decimal(order.approved_snapshot["lines"][0]["kg_per_piece"]) == 2
    assert order.approved_snapshot["export_profiles"][str(profile.id)]["palletization"]["enabled"] is True
    product.kg_per_piece = 10
    profile.palletization = {"enabled": False}
    db_session.commit()
    assert client.get(f"/api/v1/orders/{order.id}").json()["pallet_profile_ids"] == [profile.id]
    content, media, filename = ExportEngine.export_order(db_session, order.id, profile.id)
    assert media == "application/zip" and filename.endswith(".zip")
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        assert len(zf.namelist()) == 2
    db_session.refresh(order, ["export_records"])
    assert order.export_records[-1].content_hash
    late = ExportProfileService.create_profile(db_session, ExportProfileCreate(
        company_id=company.id, name="Late pallet", format="order_sheet", palletization=cfg))
    with pytest.raises(OrderExportError, match="profile frozen"):
        ExportEngine.export_order(db_session, order.id, late.id)
    historical = copy.deepcopy(order.approved_snapshot)
    historical["lines"][0].pop("kg_per_piece")
    order.approved_snapshot = historical
    db_session.commit()
    with pytest.raises(OrderExportError, match="SKU A.*no frozen trusted"):
        ExportEngine.export_order(db_session, order.id, profile.id)


def test_weight_limited_profile_rejects_approval_until_trusted_weight_is_added(client, db_session):
    company = Company(name="Approval weights")
    db_session.add(company); db_session.flush()
    customer = Customer(company_id=company.id, customer_code="C", customer_name="Buyer")
    product = Product(company_id=company.id, sku="MISSING", description="No weight", unit="piece")
    db_session.add_all([customer, product]); db_session.flush()
    profile = ExportProfileService.create_profile(db_session, ExportProfileCreate(
        company_id=company.id, name="Weight limit", format="order_sheet",
        palletization=config(max_weight_kg=10)))
    order = Order(company_id=company.id, customer_id=customer.id, order_number="O-weight", raw_input="1 MISSING", status="pending_review")
    db_session.add(order); db_session.flush()
    order.lines.append(OrderLine(company_id=company.id, line_number=1, original_text="1 MISSING", product_phrase="MISSING",
        requested_quantity=1, requested_unit="piece", final_quantity=1, final_unit="piece", matched_product_id=product.id,
        final_sku="MISSING", confidence_score=1, status="confirmed"))
    db_session.commit()

    rejected = client.post(f"/api/v1/orders/{order.id}/approve")
    assert rejected.status_code == 400
    assert "Weight limit" in rejected.json()["detail"] and "SKU MISSING" in rejected.json()["detail"]
    db_session.refresh(order)
    assert order.status == "pending_review" and order.approved_snapshot is None
    assert client.get(f"/api/v1/orders/{order.id}").json()["pallet_profile_ids"] == []

    product.kg_per_piece = Decimal("0.75")
    db_session.commit()
    OrderWorkflowService.approve_order(db_session, order.id)
    assert Decimal(order.approved_snapshot["lines"][0]["kg_per_piece"]) == Decimal("0.75")
    assert client.get(f"/api/v1/orders/{order.id}").json()["pallet_profile_ids"] == [profile.id]


def test_ai_pallet_proposal_requires_resolution_and_apply(client, db_session):
    company = Company(name="AI pallet")
    db_session.add(company); db_session.flush()
    product = Product(company_id=company.id, sku="1001", description="A")
    other = Company(name="Other")
    db_session.add_all([product, other]); db_session.flush()
    foreign = Product(company_id=other.id, sku="1002", description="B")
    db_session.add(foreign); db_session.commit()
    profile = ExportProfileService.create_profile(db_session, ExportProfileCreate(
        company_id=company.id, name="Sheet", format="order_sheet"))
    response = client.post(f"/api/v1/companies/{company.id}/rules/analyze", json={
        "description": "Codes 1001 and 1002 go on one pallet. Maximum 360 kg and 16 Excel rows including gifts."})
    assert response.status_code == 200, response.text
    proposal = response.json()
    assert proposal["export_patch"]["palletization"]["automatic_pallets"]["max_rows"] == 16
    refs = proposal["export_patch"]["palletization"]["dedicated_groups"][0]["products"]
    assert refs[0]["product_id"] == product.id and refs[1]["product_id"] is None
    assert db_session.get(ExportProfile, profile.id).palletization["enabled"] is False
    proposal["export_patch"]["palletization"]["dedicated_groups"][0]["products"] = [refs[0]]
    applied = client.post(f"/api/v1/companies/{company.id}/rules/apply-proposal", json={"proposal": proposal})
    assert applied.status_code == 200, applied.text
    assert db_session.get(ExportProfile, profile.id).palletization["dedicated_groups"][0]["product_ids"] == [product.id]


@pytest.mark.parametrize("description,mode,layout", [
    ("Maximum 20 product lines per pallet. Gift line doesn't count. Each pallet in a different sheet.",
     "logical_product_lines", "multi_sheet_workbook"),
    ("Maximum 500 kg and 25 exported rows per pallet. Different Excel file for every pallet.",
     "output_rows", "separate_workbook_per_pallet"),
])
def test_mock_ai_proposes_typed_row_mode_and_layout(description, mode, layout):
    proposal = _mock_analysis(description).export_patch.palletization
    assert proposal.automatic_pallets.row_count_mode == mode
    assert proposal.output.layout == layout


def test_product_excel_gross_kg_per_piece_reimport(db_session):
    company = Company(name="Weights")
    db_session.add(company); db_session.commit()
    wb = openpyxl.Workbook()
    wb.active.append(["SKU", "Description", "Μικτό βάρος"])
    wb.active.append(["W-1", "Widget", "0,75"])
    output = io.BytesIO(); wb.save(output)
    file_bytes = output.getvalue()
    preview = MasterDataService.preview_excel(file_bytes, "products")
    assert preview.suggested_mapping["kg_per_piece"] == "Μικτό βάρος"
    mapping = {"sku": "SKU", "description": "Description", "kg_per_piece": "Μικτό βάρος"}
    first = MasterDataService.import_products(db_session, company.id, file_bytes, mapping)
    assert first.imported == 1 and first.errors == 0
    product = db_session.query(Product).filter_by(company_id=company.id, sku="W-1").one()
    assert product.kg_per_piece == Decimal("0.75")
    wb.active["C2"] = "0,80"
    output = io.BytesIO(); wb.save(output)
    second = MasterDataService.import_products(db_session, company.id, output.getvalue(), mapping)
    assert second.imported == 1 and second.errors == 0
    db_session.refresh(product)
    assert product.kg_per_piece == Decimal("0.80")


def test_weight_edit_cannot_cross_company(client, db_session):
    a, b = Company(name="Weight A"), Company(name="Weight B")
    db_session.add_all([a, b]); db_session.flush()
    product = Product(company_id=a.id, sku="A", description="A")
    db_session.add(product); db_session.commit()
    response = client.put(f"/api/v1/products/{product.id}/physical-weight", json={"company_id": b.id, "kg_per_piece": 5})
    assert response.status_code == 404
    assert db_session.get(Product, product.id).kg_per_piece is None
