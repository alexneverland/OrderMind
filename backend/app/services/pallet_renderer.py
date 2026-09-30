"""Render a completed PalletPlan without performing allocation."""
import io
import zipfile

import openpyxl

from backend.app.schemas.pallet import PalletConfig
from backend.app.services.export_registry import sanitize_formula_injection

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ZIP_MIME = "application/zip"
HEADERS = ["SKU", "Description", "Bonus", "Quantity"]


def _append_pallet(ws, pallet, *, title: bool, header: bool):
    if title:
        label = f"PALLET {pallet['pallet_number']}"
        if pallet["group_name"]:
            label += f" · {pallet['group_name']}"
        ws.append([sanitize_formula_injection(label)])
    if header:
        ws.append(HEADERS)
    for item in pallet["items"]:
        for row in item["rows"]:
            ws.append(row)


def _bytes(wb) -> bytes:
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def render_pallet_plan(order_id: int, plan: dict, config: PalletConfig, include_header: bool) -> tuple[bytes, str, str]:
    pallets = plan["pallets"]
    layout, output = config.output.layout, config.output
    if layout == "single_sheet_sections":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Order"
        for index, pallet in enumerate(pallets):
            if index:
                for _ in range(output.blank_rows_between_pallets):
                    ws.append([])
            _append_pallet(ws, pallet, title=output.show_pallet_title,
                           header=(include_header and index == 0) or output.repeat_headers)
        return _bytes(wb), XLSX_MIME, f"ordermind_order_{order_id}_pallets.xlsx"
    if layout == "multi_sheet_workbook":
        wb = openpyxl.Workbook()
        for index, pallet in enumerate(pallets):
            ws = wb.active if index == 0 else wb.create_sheet()
            ws.title = f"Pallet {index + 1}"[:31]
            _append_pallet(ws, pallet, title=output.show_pallet_title, header=include_header)
        return _bytes(wb), XLSX_MIME, f"ordermind_order_{order_id}_pallets.xlsx"
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for index, pallet in enumerate(pallets, 1):
            wb = openpyxl.Workbook()
            wb.active.title = "Order"
            _append_pallet(wb.active, pallet, title=output.show_pallet_title, header=include_header)
            info = zipfile.ZipInfo(f"ordermind_order_{order_id}_pallet_{index:02d}.xlsx", (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            zf.writestr(info, _bytes(wb))
    return archive.getvalue(), ZIP_MIME, f"ordermind_order_{order_id}_pallets.zip"
