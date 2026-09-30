"""Extract reviewable text from customer order attachments without storing binaries."""

import asyncio
from io import BytesIO
from pathlib import Path

from backend.app.config import settings
from backend.app.core.office_archive import validate_office_archive, OfficeArchiveError
from backend.app.ai.prompt_boundaries import OCR_SYSTEM_PROMPT


SUPPORTED_EXTENSIONS = {".txt", ".csv", ".xlsx", ".docx", ".pdf", ".png", ".jpg", ".jpeg", ".webp"}
IMAGE_MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}


class OrderFileError(ValueError):
    pass


def _bounded_text(text: str) -> str:
    text = text.strip()
    if not text:
        raise OrderFileError("No readable order text was found in this file")
    if len(text) > settings.MAX_RAW_ORDER_TEXT_SIZE:
        raise OrderFileError("Extracted order text exceeds the 50,000 character limit")
    return text


def _plain_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1253"):
        try:
            return _bounded_text(content.decode(encoding))
        except UnicodeDecodeError:
            continue
    raise OrderFileError("Text file must use UTF-8 or Windows Greek encoding")


def _spreadsheet(content: bytes) -> str:
    validate_office_archive(content)
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) > 20:
            raise OrderFileError("Spreadsheet has more than 20 sheets")
        lines = []
        total_chars = 0
        for sheet in workbook:
            lines.append(f"Sheet: {sheet.title}")
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if row_number > 1000:
                    raise OrderFileError("Spreadsheet has more than 1,000 rows per sheet")
                if len(row) > 30:
                    raise OrderFileError("Spreadsheet has more than 30 columns")
                cells = [str(value).strip() if value is not None else "" for value in row]
                if any(cells):
                    # Preserve empty cells so the AI can infer associations from
                    # column position even when the sheet layout varies.
                    line = " | ".join(cells)
                    lines.append(line)
                    total_chars += len(line)
                if total_chars > settings.MAX_RAW_ORDER_TEXT_SIZE:
                    raise OrderFileError("Extracted order text exceeds the 50,000 character limit")
        return _bounded_text("\n".join(lines))
    finally:
        workbook.close()


def _document(content: bytes) -> str:
    validate_office_archive(content)
    from docx import Document

    document = Document(BytesIO(content))
    # XML document order is preserved for paragraphs and tables.
    lines = []
    for block in document.iter_inner_content():
        if hasattr(block, "rows"):
            for row in block.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        else:
            lines.append(block.text.strip())
    return _bounded_text("\n".join(line for line in lines if line))


def _pdf(content: bytes) -> tuple[str, bool]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content), strict=False)
    if len(reader.pages) > 20:
        raise OrderFileError("PDF has more than 20 pages")
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    if not pages or any(not page for page in pages):
        return "", True  # A scanned page needs OCR, including in a mixed PDF.
    return _bounded_text("\n".join(pages)), False


OCR_PROMPT = "Transcribe all visible printed or handwritten text from this customer order in reading order. Preserve quantities, units and product wording exactly. If any word or number cannot be read reliably, write [UNCLEAR] in its place. Do not infer missing text, add products, or summarize. Return only the plain-text transcription."


async def _transcribe(content: bytes, mime_type: str) -> str:
    provider_name = settings.AI_PROVIDER.strip().lower()
    if provider_name == "mock":
        raise OrderFileError("Photo or scanned PDF text needs an AI provider. Choose Gemini, OpenAI, Claude or Vertex in Settings")
    from backend.app.ai.factory import get_ai_provider
    from google.genai import types

    provider = get_ai_provider()
    if provider_name in {"gemini", "vertex"}:
        client = provider._get_client()
        try:
            for attempt in range(3):
                try:
                    response = await client.aio.models.generate_content(
                        model=provider.model_name,
                        contents=[OCR_PROMPT, types.Part.from_bytes(data=content, mime_type=mime_type)],
                        config=types.GenerateContentConfig(temperature=0, system_instruction=OCR_SYSTEM_PROMPT),
                    )
                    return _bounded_text(response.text or "")
                except Exception as exc:
                    if attempt == 2 or getattr(exc, "code", None) not in {429, 500, 502, 503, 504}:
                        raise
                    await asyncio.sleep(2 ** attempt)
        finally:
            await client.aio.aclose()
            client.close()
    return _bounded_text(await provider.transcribe_file(content, mime_type, OCR_PROMPT))


async def extract_order_file(filename: str, content: bytes) -> tuple[str, str]:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise OrderFileError("Unsupported order file. Use TXT, CSV, XLSX, DOCX, PDF, PNG, JPG or WebP")
    if not content:
        raise OrderFileError("The uploaded file is empty")
    if len(content) > settings.MAX_ORDER_UPLOAD_SIZE_BYTES:
        raise OrderFileError("Order file exceeds the 10 MB limit")
    try:
        if extension in {".txt", ".csv"}:
            return _plain_text(content), "text"
        if extension == ".xlsx":
            return _spreadsheet(content), "spreadsheet"
        if extension == ".docx":
            return _document(content), "document"
        if extension == ".pdf":
            if not content.startswith(b"%PDF"):
                raise OrderFileError("This is not a valid PDF file")
            text, needs_ocr = _pdf(content)
            if needs_ocr:
                return await _transcribe(content, "application/pdf"), "ocr"
            return text, "pdf"
        if extension == ".png" and not content.startswith(b"\x89PNG\r\n\x1a\n"):
            raise OrderFileError("This is not a valid PNG image")
        if extension in {".jpg", ".jpeg"} and not content.startswith(b"\xff\xd8\xff"):
            raise OrderFileError("This is not a valid JPEG image")
        if extension == ".webp" and not (content.startswith(b"RIFF") and content[8:12] == b"WEBP"):
            raise OrderFileError("This is not a valid WebP image")
        return await _transcribe(content, IMAGE_MIME[extension]), "ocr"
    except OrderFileError:
        raise
    except OfficeArchiveError as exc:
        raise OrderFileError(str(exc)) from exc
    except Exception as exc:
        raise OrderFileError("Could not read this order file; check that it is valid and unprotected") from exc
